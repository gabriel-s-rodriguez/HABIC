# -*- coding: utf-8 -*-
"""Pruebas de HABIC. Sin frameworks, solo asserts. Ejecutar con: py test_habic.py

No sale a internet ni gasta cuota: las llamadas a Claude van simuladas.
Las exportaciones escriben en una carpeta temporal, nunca en resultados/.

Las pruebas de exportación son de CARACTERIZACIÓN: la primera vez guardan una
foto de las hojas generadas en referencia_export.json, y a partir de ahí
comparan celda a celda contra esa foto. Sirven para refactorizar sin miedo: si
un cambio altera el Excel aunque sea en una celda, salta aquí.
"""
import json
import subprocess
import sys
import tempfile
import tkinter as tk
from datetime import datetime
from pathlib import Path

import openpyxl

import habic

# La consola de Windows llega en cp1252 y se atraganta con los acentos
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

AQUI       = Path(__file__).parent
REFERENCIA = AQUI / "referencia_export.json"
TMP        = Path(tempfile.mkdtemp(prefix="habic_pruebas_"))

# Las exportaciones y las plantillas van a la carpeta temporal
habic.RESULTS_DIR = TMP

# Ningún diálogo puede bloquear la ejecución
habic.messagebox.showerror   = lambda *a, **k: None
habic.messagebox.showinfo    = lambda *a, **k: None
habic.messagebox.showwarning = lambda *a, **k: None
habic.messagebox.askyesno    = lambda *a, **k: True

_hechas = [0]


def check(cond, msg):
    _hechas[0] += 1
    if not cond:
        raise AssertionError(msg)


def seccion(titulo):
    print(f"\n── {titulo} " + "─" * max(0, 58 - len(titulo)))


# ══════════════════════════════════════════════════════════════════════════
seccion("Cálculo local (RAM, lote, contexto)")

check(habic.recommended_batch(None) == 5, "sin RAM detectada → lote 5")
check(habic.recommended_batch(4) == 1, "4 GB → lote 1")
check(habic.recommended_batch(8) == 3, "8 GB → lote 3")
check(habic.recommended_batch(16) == 5, "16 GB → lote 5")
check(habic.recommended_batch(64) == 8, "64 GB → lote 8")
check(habic.recommended_max_ctx(None) == 4096, "sin RAM → ctx 4096")
check(habic.recommended_max_ctx(4) == 2048, "4 GB → ctx 2048")
check(habic.recommended_max_ctx(64) == 16384, "64 GB → ctx 16384")
check(habic.choose_num_ctx(1000, 8192) == (2048, True), "cabe en la ventana chica")
check(habic.choose_num_ctx(5000, 8192) == (8192, True), "sube a la siguiente")
check(habic.choose_num_ctx(9000, 8192) == (8192, False),
      "no cabe → fits=False para poder avisar")
ram = habic.detect_ram_gb()
check(ram is None or ram in (4, 8, 16, 32, 64, 128) or isinstance(ram, int),
      f"RAM detectada rara: {ram}")
print(f"RAM de este equipo: {ram} GB")


# ══════════════════════════════════════════════════════════════════════════
seccion("Codebook: carga y validación")

CB_DIR = TMP / "codebooks"
CB_DIR.mkdir()


def escribir_codebook(nombre, filas, encabezados=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(encabezados or ["CODIGO", "Nombre", "Descripcion", "Criterios",
                              "Ejemplos_Positivos", "Ejemplos_Exclusion"])
    for f in filas:
        ws.append(f)
    ruta = CB_DIR / nombre
    wb.save(ruta)
    return ruta


bueno = escribir_codebook("bueno.xlsx", [
    ["CRIT", "Crítica", "Critica al gobierno", "Señala fallos", "ej sí", "ej no"],
    ["MOVI", "Movilización", "Convoca", "Llama a marchar", "ej sí", "ej no"],
])
escribir_codebook("otro.xlsx", [["A", "a", "d", "c", "", ""]])

cats = habic.load_codebook(bueno)
check(len(cats) == 2, f"debía cargar 2 categorías, cargó {len(cats)}")
check(cats[0]["codigo"] == "CRIT",
      "los encabezados en MAYÚSCULAS deben normalizarse a minúsculas")
check(cats[1]["nombre"] == "Movilización", "segunda categoría mal leída")

habic.CODEBOOK_DIR = CB_DIR
encontrados = habic.scan_codebooks()
check([n for n, _ in encontrados] == ["bueno", "otro"],
      f"scan_codebooks debe devolver ordenado: {encontrados}")

err, avisos = habic.validate_codebook(cats, "bueno.xlsx")
check(not err, f"el codebook bueno no debe dar errores: {err}")

err, _ = habic.validate_codebook([], "vacio.xlsx")
check(err, "un codebook vacío debe dar error crítico")

err, _ = habic.validate_codebook(
    [{"codigo": "A", "nombre": "a"}], "corto.xlsx")
check(err, "faltan columnas obligatorias → error crítico")

dup = [dict(cats[0]), dict(cats[0])]
err, _ = habic.validate_codebook(dup, "dup.xlsx")
check(any("uplicad" in e for e in err), f"códigos duplicados → error: {err}")

sin_ej = [{k: v for k, v in c.items()
           if k not in ("ejemplos_positivos", "ejemplos_exclusion")}
          for c in cats]
err, avisos = habic.validate_codebook(sin_ej, "sinej.xlsx")
check(not err and avisos,
      "sin ejemplos: advertencia, no error (el análisis puede correr)")


# ══════════════════════════════════════════════════════════════════════════
seccion("Prompt, parseo y matriz")

ITEMS = [{"id": "t1", "texto": "El gobierno incumplió"},
         {"id": "t2", "texto": "Mañana a la plaza"}]

prompt = habic.build_batch_prompt(ITEMS, cats)
for trozo in ("CRIT", "MOVI", "Señala fallos", "t1", "t2",
              "El gobierno incumplió", "razon"):
    check(trozo in prompt, f"el prompt debe incluir '{trozo}'")
check(prompt.count("[ID:") == 2, "un bloque por texto")

check(habic._try_parse_json('{"a": 1,}') == {"a": 1}, "coma colgante")
check(habic._try_parse_json('bla bla {"a": 2} fin') == {"a": 2},
      "JSON rodeado de prosa")
check(habic._try_parse_json("no hay json") is None, "sin JSON → None")

bien = json.dumps({
    "t1": {"CRIT": {"valor": 1, "razon": "cita"},
           "MOVI": {"valor": 0, "razon": "no"}},
    "t2": {"CRIT": {"valor": 0, "razon": "no"},
           "MOVI": {"valor": "NC", "razon": "ambiguo"}},
})
res = habic.parse_batch_response(bien, ITEMS, cats)
check(res["t1"]["CRIT"]["valor"] == 1, "valor 1 mal parseado")
check(res["t2"]["MOVI"]["valor"] == "NC", "NC debe conservarse como NC")

res_malo = habic.parse_batch_response("no es json", ITEMS, cats)
check(res_malo["t1"]["CRIT"]["valor"] is None,
      "sin respuesta → None (dato faltante), NUNCA 0")

vacia = habic.parse_batch_response(
    json.dumps({"t1": {"CRIT": {"valor": 1, "razon": "  "}},
                "t2": {}}), ITEMS, cats)
check("no proporcionó justificación" in vacia["t1"]["CRIT"]["razon"],
      "razón vacía debe rellenarse con el aviso")

check(habic.celda_matriz(res["t1"], "CRIT") == 1, "matriz: 1")
check(habic.celda_matriz(res["t1"], "MOVI") == 0, "matriz: 0")
check(habic.celda_matriz(res["t2"], "MOVI") == "N/C", "matriz: N/C")
check(habic.celda_matriz(res_malo["t1"], "CRIT") is None, "matriz: vacía")


# ══════════════════════════════════════════════════════════════════════════
seccion("Motor Claude (llamadas simuladas)")


class FalsoProceso:
    def __init__(self, stdout="", stderr="", returncode=0):
        self.stdout, self.stderr, self.returncode = stdout, stderr, returncode


def simular(respuesta):
    ultimo = {}

    def falso_run(cmd, **kw):
        ultimo["cmd"] = cmd
        if isinstance(respuesta, Exception):
            raise respuesta
        return respuesta

    subprocess.run = falso_run
    return ultimo


def envoltorio(texto, is_error=False, coste=0.0):
    return json.dumps({"result": texto, "is_error": is_error,
                       "total_cost_usd": coste})


run_original = subprocess.run
try:
    ultimo = simular(FalsoProceso(stdout=envoltorio(bien)))
    crudo, coste = habic.claude_generate("prompt", "sonnet", "low")
    check(coste == 0.0, "coste 0 cuando el envoltorio dice 0")
    check(habic.parse_batch_response(crudo, ITEMS, cats)["t1"]["CRIT"]["valor"] == 1,
          "el circuito CLI → parser no entrega la codificación")

    # Los tres flags de aislamiento no son opcionales: sin ellos el CLAUDE.md
    # del usuario y sus skills contaminarían el contexto del codificador.
    cmd = ultimo["cmd"]
    for flag in ("--no-session-persistence", "--setting-sources",
                 "--allowedTools", "--output-format"):
        check(flag in cmd, f"falta el flag {flag}")
    check(cmd[cmd.index("--model") + 1] == "sonnet", "modelo mal pasado")
    check(cmd[cmd.index("--effort") + 1] == "low", "esfuerzo mal pasado")

    for msg in ("Claude usage limit reached. Resets at 4pm",
                "Error: not authenticated", "model not found"):
        simular(FalsoProceso(stderr=msg, returncode=1))
        try:
            habic.claude_generate("p", "sonnet", "low")
            raise AssertionError(f"debió parar la ronda: {msg}")
        except habic.ClaudeTerminal:
            _hechas[0] += 1

    # El CLI mete el error dentro del envoltorio y sale con código != 0
    simular(FalsoProceso(
        stdout=envoltorio("Failed to authenticate: OAuth session expired and "
                          "could not be refreshed", is_error=True),
        returncode=1))
    try:
        habic.claude_generate("p", "sonnet", "low")
        raise AssertionError("debió lanzar ClaudeTerminal")
    except habic.ClaudeTerminal as e:
        check("OAuth" in str(e), f"el mensaje debe ser legible, no el JSON: {e}")

    for fallo in (FalsoProceso(stderr="connection reset by peer", returncode=1),
                  subprocess.TimeoutExpired(cmd="claude", timeout=1)):
        simular(fallo)
        try:
            habic.claude_generate("p", "sonnet", "low")
            raise AssertionError("debió lanzar RuntimeError")
        except habic.ClaudeTerminal:
            raise AssertionError("un fallo pasajero no debe parar la ronda")
        except RuntimeError:
            _hechas[0] += 1

    simular(FalsoProceso(stdout=envoltorio("usage limit reached", is_error=True)))
    try:
        habic.claude_generate("p", "sonnet", "low")
        raise AssertionError("is_error con código 0 debe parar la ronda")
    except habic.ClaudeTerminal:
        _hechas[0] += 1

    # JSON que no es el envoltorio del CLI: se pasa tal cual al parser
    simular(FalsoProceso(stdout=bien))
    check(habic.parse_batch_response(
        habic.claude_generate("p", "sonnet", "low")[0], ITEMS, cats
    )["t1"]["CRIT"]["valor"] == 1, "respuesta sin envoltorio perdida")

    simular(FalsoProceso(stdout=envoltorio(bien, coste=0.0412)))
    _, coste = habic.claude_generate("p", "sonnet", "low")
    check(abs(coste - 0.0412) < 1e-9, f"coste mal leído: {coste}")
finally:
    subprocess.run = run_original


# ══════════════════════════════════════════════════════════════════════════
seccion("Fiabilidad: kappa y muestra de validación")

check(habic.kappa_for_pair([]) == (None, None, None, 0), "sin pares → sin κ")
k, po, pe, n = habic.kappa_for_pair([(1, 1), (0, 0), (1, 1), (0, 0)])
check(k == 1.0 and po == 1.0 and n == 4, f"acuerdo total → κ=1: {k}, {po}")
k, po, _, _ = habic.kappa_for_pair([(1, 0), (0, 1), (1, 0), (0, 1)])
check(po == 0.0 and k < 0, f"desacuerdo total → κ negativo: {k}")
k, _, pe, _ = habic.kappa_for_pair([(1, 1), (1, 1)])
check(pe == 1.0 and k == 1.0, "todos unos: pe=1 y acuerdo → κ=1, sin dividir por 0")
k, _, pe, _ = habic.kappa_for_pair([(1, 1), (1, 0)])
check(k is not None, "no debe reventar con prevalencia extrema")

check(habic.lacy_riffe_n(1) == 1, "N=1 → n=1")
check(habic.lacy_riffe_n(100) <= 100, "la muestra nunca supera el corpus")
check(habic.lacy_riffe_n(1000) < 1000, "con N grande la muestra es menor")
check(habic.kappa_interpretation(None) == "—", "sin κ → guion")
check(habic.kappa_interpretation(-0.1) == "Sin acuerdo", "κ negativo")
check(habic.kappa_interpretation(0.1) == "Pobre", "κ 0,1")
check(habic.kappa_interpretation(0.5) == "Moderado", "κ 0,5")
check(habic.kappa_interpretation(0.85) == "Casi perfecto", "κ 0,85")


def escribir_resultado(nombre, filas):
    """Un .xlsx con el formato que produce HABIC (hoja Corpus)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Corpus"
    ws.append(["id", "texto", "CRIT", "MOVI"])
    for f in filas:
        ws.append(f)
    ruta = TMP / nombre
    wb.save(ruta)
    return ruta


ruta_a = escribir_resultado("res_a.xlsx", [
    ["1", "uno", 1, 0], ["2", "dos", 0, 1], ["3", "tres", 1, None]])
data, codes, meta = habic.load_result_file(ruta_a)
# load_result_file normaliza los encabezados a minúsculas: los códigos vuelven
# en minúscula aunque en el archivo estén en mayúscula.
check(codes == ["crit", "movi"], f"columnas de categoría mal detectadas: {codes}")
check(data["1"]["crit"] == 1, "valor mal leído")
check(data["3"]["movi"] is None, "celda vacía debe ser None, no 0")
check(meta["2"]["texto"] == "dos", "columna meta perdida")

# Solo encabezados: carga sin datos, no es un error (el κ lo detecta después)
solo_cab, _, _ = habic.load_result_file(escribir_resultado("solo_cab.xlsx", []))
check(solo_cab == {}, f"solo encabezados → sin filas: {solo_cab}")

for nombre, filas, cabecera in (
        ("hoja_vacia.xlsx", None, None),
        ("sin_id.xlsx", [["uno", 1]], ["texto", "CRIT"]),
        ("sin_cats.xlsx", [["1", "uno"]], ["id", "texto"])):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Corpus"
    if cabecera:
        ws.append(cabecera)
        for f in filas:
            ws.append(f)
    ruta = TMP / nombre
    wb.save(ruta)
    try:
        habic.load_result_file(ruta)
        raise AssertionError(f"{nombre} debía dar error")
    except ValueError:
        _hechas[0] += 1


# ══════════════════════════════════════════════════════════════════════════
seccion("Interfaz: construcción, motores y diálogos")

root = tk.Tk()
root.withdraw()
app = habic.HABICApp(root)
check(app is not None, "la aplicación no se construyó")

app.motor_var.set("claude")
app._aplicar_motor()
check(list(app.model_combo["values"]) == list(habic.CLAUDE_MODELS),
      "al pulsar Claude deben salir sus modelos")
check(app.model_var.get() in habic.CLAUDE_MODELS, "modelo de Claude no elegido")
check(str(app.btn_compatibles["state"]) == "disabled",
      "los botones de Ollama deben apagarse con Claude")
check(app.effort_combo.winfo_manager() == "pack", "el esfuerzo debe verse")

app.refresh_models()
check(list(app.model_combo["values"]) == list(habic.CLAUDE_MODELS),
      "refrescar Ollama no puede pisar la lista de Claude")

app.motor_var.set("ollama")
app._aplicar_motor()
check(list(app.model_combo["values"]) == app._ollama_models,
      "al volver a Ollama deben salir los modelos locales")
check(app.effort_combo.winfo_manager() == "", "el esfuerzo debe ocultarse")
check(str(app.btn_compatibles["state"]) == "normal",
      "los botones de Ollama deben reactivarse")

check(app._nombre_motores([("ollama", "m", None)]) == "Ollama", "nombre motor")
check(app._nombre_motores([("claude", "sonnet", "low")]) == "Claude", "nombre motor")
check(app._nombre_motores([("ollama", "m", None),
                           ("claude", "s", "low")]) == "Ollama y Claude",
      "mezcla de motores mal nombrada")
check(app._model_available("claude", "") is False, "sin modelo → no disponible")
check(app._model_available("claude", "sonnet") is bool(habic.CLAUDE_BIN),
      "disponibilidad de Claude según el ejecutable")

texto = app._texto_estimacion(800)
check("800 textos" in texto, f"la estimación debe decir cuántos textos: {texto}")
check("tres veces" in texto, "con 800 textos avisa tres veces")
check("una vez" in app._texto_estimacion(50), "con 50 textos avisa una vez")

# Fechas
app.cuenta_var.set("@alguien")
app.desde_var.set("01/10/2023")
app.hasta_var.set("15/10/2023")
ok_inputs = app.validate_inputs()
check(ok_inputs and ok_inputs[0] == "alguien", "la arroba debe quitarse")
app.hasta_var.set("01/09/2023")
check(app.validate_inputs() is None, "hasta anterior a desde → inválido")
app.hasta_var.set("32/13/2023")
check(app.validate_inputs() is None, "fecha imposible → inválido")
app.cuenta_var.set("")
app.desde_var.set("01/10/2023"); app.hasta_var.set("15/10/2023")
check(app.validate_inputs() is None, "sin cuenta → inválido")

entrada = tk.Entry(root)
var = tk.StringVar(value="01102023")
app.autoformat_date(var, entrada)
check(var.get() == "01/10/2023", f"autoformato de fecha: {var.get()}")

# Plantillas descargables, con el diálogo de archivo sustituido
destinos = {}
for nombre, metodo in (("codebook", app._cb_download_template),
                       ("corpus", app._xl_download_template),
                       ("kappa", app._kappa_download_template)):
    destino = TMP / f"plantilla_{nombre}.xlsx"
    habic.filedialog.asksaveasfilename = (
        lambda *a, _d=str(destino), **k: _d)
    metodo()
    check(destino.exists(), f"no se generó la plantilla de {nombre}")
    wb = openpyxl.load_workbook(destino)
    check(wb.active.max_row >= 1, f"plantilla de {nombre} vacía")
    destinos[nombre] = destino


# ══════════════════════════════════════════════════════════════════════════
seccion("Codificación: completa, cancelada, parada y segundo pase")

CATS1  = [{"codigo": "X", "nombre": "x"}]
I40    = [{"id": f"i{n}", "texto": f"texto {n}"} for n in range(40)]
I200   = [{"id": f"j{n}", "texto": f"t {n}"} for n in range(200)]


def motor_falso(items, valor=1, razon="r"):
    return lambda motor, model, effort, prompt, ni, nc, fn, pre: json.dumps(
        {i["id"]: {"X": {"valor": valor, "razon": razon}} for i in items})


app.prefs["control_activo"] = True
app._cancel_event.clear()
app._ronda_coste = 0.0
app._generar_raw = motor_falso(I40)

preguntas = []
app._confirmar_control = lambda hechos, *a: (preguntas.append(hechos), True)[1]
res, motivo = app._code_items(I40, CATS1, "claude", "sonnet", "low",
                              lambda m: None, lambda m: None, app.xl_progress)
check(motivo is None, f"ronda completa: motivo {motivo}")
check(len(res) == 40 and all(r["coding"]["X"]["valor"] == 1 for r in res),
      "la ronda completa debe codificar los 40")
check(len(preguntas) == 1, f"menos de 100 textos → 1 aviso, hubo {len(preguntas)}")
check(20 <= preguntas[0] < 40, f"el aviso va a la mitad: {preguntas}")

app._generar_raw = motor_falso(I200)
preguntas.clear()
app._code_items(I200, CATS1, "claude", "sonnet", "low",
                lambda m: None, lambda m: None, app.xl_progress)
check(len(preguntas) == 3, f"200 textos → 3 avisos, hubo {len(preguntas)}")

app.prefs["control_activo"] = False
preguntas.clear()
app._code_items(I200, CATS1, "claude", "sonnet", "low",
                lambda m: None, lambda m: None, app.xl_progress)
check(not preguntas, f"desactivado y preguntó {len(preguntas)} veces")

app.prefs["control_activo"] = True
preguntas.clear()
app._code_items(I200, CATS1, "ollama", "llama3.1:8b", None,
                lambda m: None, lambda m: None, app.xl_progress)
check(not preguntas, "Ollama no consume cuota: no debe preguntar")

# Parar en un control conserva TODAS las filas
app._generar_raw = motor_falso(I40)
app._confirmar_control = lambda *a: False
res, motivo = app._code_items(I40, CATS1, "claude", "sonnet", "low",
                              lambda m: None, lambda m: None, app.xl_progress)
check(motivo == "parado", f"motivo esperado 'parado': {motivo}")
check(len(res) == 40, f"el parcial perdió filas: {len(res)}")
hechos = [r for r in res if r["coding"]["X"]["valor"] == 1]
falta  = [r for r in res if r["coding"]["X"]["valor"] is None]
check(hechos and falta, f"parcial mal partido: {len(hechos)}/{len(falta)}")
check("Pendiente" in falta[0]["coding"]["X"]["razon"], "razón del pendiente")
check(habic.celda_matriz(falta[0]["coding"], "X") is None,
      "un pendiente va como celda vacía, nunca como 0")

# Cancelar no devuelve resultados
app._confirmar_control = lambda *a: True
app._cancel_event.set()
res, motivo = app._code_items(I40, CATS1, "claude", "sonnet", "low",
                              lambda m: None, lambda m: None, app.xl_progress)
check(motivo == "cancelado" and res is None, "cancelar no exporta nada")
app._cancel_event.clear()

# Segundo pase: los NC se repreguntan uno a uno y pueden resolverse
llamadas = [0]


def motor_nc(motor, model, effort, prompt, ni, nc, fn, pre):
    llamadas[0] += 1
    valor = "NC" if llamadas[0] == 1 else 1
    return json.dumps({i["id"]: {"X": {"valor": valor, "razon": "r"}}
                       for i in I40[:3]})


app._generar_raw = motor_nc
app._confirmar_ui = lambda *a: True      # aceptar reprocesar los dudosos
res, motivo = app._code_items(I40[:3], CATS1, "claude", "sonnet", "low",
                              lambda m: None, lambda m: None, app.xl_progress)
check(all(r["coding"]["X"]["valor"] == 1 for r in res),
      "el segundo pase debe resolver los NC que el modelo sí decide")
check(llamadas[0] > 1, "el segundo pase ni siquiera llamó al modelo")


# ══════════════════════════════════════════════════════════════════════════
seccion("Fin de ronda: qué botón queda encendido en cada pestaña")


def estado(b):
    return str(b["state"])


def fin_de_ronda(quien, con_datos_en_la_otra):
    """Deja la app como si acabara de terminar una codificación."""
    app.tweets       = [{"id": "1"}] if quien == "corpus" and con_datos_en_la_otra else []
    app.corpus_items = [{"id": "1"}] if quien == "twitter" and con_datos_en_la_otra else []
    app._process_running = True
    for b in (app.btn_start, app.btn_code, app.btn_cancel,
              app.btn_xl_code, app.btn_xl_cancel):
        b.config(state=tk.DISABLED)


# Twitter termina: sus tres botones vuelven, y el de corpus solo si hay corpus
for con_corpus in (True, False):
    fin_de_ronda("twitter", con_corpus)
    app.on_twitter_all_done(["a.xlsx"], None, 0)
    root.update()
    check(app._process_running is False, "el proceso debe quedar marcado como fin")
    check(estado(app.btn_start) == "normal", "btn_start debe volver")
    check(estado(app.btn_code) == "normal", "btn_code debe volver")
    check(estado(app.btn_cancel) == "disabled", "btn_cancel debe apagarse")
    check(app.tw_count_var.get() == "", "el contador debe limpiarse")
    check("Listo!" in app.tw_status_var.get()
          or "¡Listo!" in app.tw_status_var.get(),
          f"barra de Twitter: {app.tw_status_var.get()}")
    if con_corpus:
        check(estado(app.btn_xl_code) == "normal",
              "con corpus cargado, su botón vuelve")
        check("Archivo listo" in app.xl_status_var.get(),
              f"mensaje de corpus: {app.xl_status_var.get()}")
    else:
        check(estado(app.btn_xl_code) == "disabled",
              "sin corpus cargado, su botón sigue apagado")
        check("Selecciona un archivo Excel" in app.xl_status_var.get(),
              f"mensaje de corpus: {app.xl_status_var.get()}")

# Corpus termina: el suyo vuelve, btn_start siempre, btn_code solo si hay trinos
for con_trinos in (True, False):
    fin_de_ronda("corpus", con_trinos)
    app.on_corpus_all_done(["a.xlsx"], None, 0)
    root.update()
    check(estado(app.btn_xl_code) == "normal", "btn_xl_code debe volver")
    check(estado(app.btn_xl_cancel) == "disabled", "btn_xl_cancel debe apagarse")
    check(estado(app.btn_start) == "normal", "btn_start vuelve siempre")
    check(app.xl_count_var.get() == "", "el contador debe limpiarse")
    check(app.tw_status_var.get() == "Listo para iniciar.",
          f"barra de Twitter: {app.tw_status_var.get()}")
    check(estado(app.btn_code) == ("normal" if con_trinos else "disabled"),
          "btn_code solo vuelve si hay trinos cargados")

# Cancelar deja los mismos estados y además pone la barra a cero
for quien, done, barra in (("twitter", app.on_twitter_coding_cancelled,
                            app.tw_progress),
                           ("corpus", app.on_corpus_coding_cancelled,
                            app.xl_progress)):
    fin_de_ronda(quien, True)
    barra.config(value=50)
    done()
    root.update()
    check(barra["value"] == 0, f"la barra de {quien} debe volver a cero")
    check(app._process_running is False, "cancelar cierra el proceso")
    check(estado(app.btn_start) == "normal", f"botones tras cancelar en {quien}")

# Un error al exportar no deja los botones bloqueados
fin_de_ronda("corpus", False)
app.on_corpus_all_done([], "disco lleno", 0)
root.update()
check(estado(app.btn_xl_code) == "normal",
      "tras un error de exportación los botones deben volver igual")

app.tweets = []
app.corpus_items = []


# ══════════════════════════════════════════════════════════════════════════
seccion("Arranque de una codificación en cada pestaña")

registro = {}


class HiloFalso:
    """Sustituye a threading.Thread: apunta con qué se iba a arrancar."""

    def __init__(self, target=None, args=(), daemon=None):
        registro["target"] = target
        registro["args"] = args

    def start(self):
        registro["arrancado"] = True


hilo_real = habic.threading.Thread
RONDAS = [("ollama", "llama3.1:8b", None)]
try:
    habic.threading.Thread = HiloFalso
    app._open_run_config = lambda modelo, n=0: RONDAS
    app.load_selected_codebook = lambda: CATS1
    app._ollama_listo = lambda r: True

    # Sin datos cargados no arranca nada
    app.tweets = []
    registro.clear()
    app.start_twitter_coding()
    check(not registro, "sin trinos no debe arrancar la codificación")
    app.corpus_items = []
    registro.clear()
    app.start_corpus_coding()
    check(not registro, "sin corpus no debe arrancar la codificación")

    # Si el usuario cancela el diálogo de rondas, tampoco
    app.tweets = [{"id": "1", "texto": "hola"}]
    app._open_run_config = lambda modelo, n=0: None
    registro.clear()
    app.start_twitter_coding()
    check(not registro, "cancelar el diálogo no debe arrancar nada")
    app._open_run_config = lambda modelo, n=0: RONDAS

    # Twitter arranca
    app.tweets = [{"id": "1", "texto": "hola"}, {"id": "2", "texto": "adiós"}]
    app.corpus_items = []
    registro.clear()
    app.start_twitter_coding()
    check(registro.get("arrancado"), "la codificación de Twitter no arrancó")
    args = registro["args"]
    check(args[0] is app.tweets, "el worker debe recibir los trinos")
    check(args[2] == RONDAS, "el worker debe recibir las rondas elegidas")
    check(args[3] == app.update_tw_status, "barra de estado de Twitter")
    check(args[5] is app.tw_progress, "barra de progreso de Twitter")
    check(args[6] == app.export_twitter_excel, "exportador de Twitter")
    check(args[7] == app.on_twitter_all_done, "callback de fin de Twitter")
    check(app._process_running is True, "debe quedar marcado como en curso")
    check(estado(app.btn_start) == "disabled", "btn_start se apaga al arrancar")
    check(estado(app.btn_code) == "disabled", "btn_code se apaga al arrancar")
    check(estado(app.btn_xl_code) == "disabled", "la otra pestaña se bloquea")
    check(estado(app.btn_cancel) == "normal", "el botón de cancelar se enciende")
    check("Hay un proceso activo en la pestaña Twitter" in app.xl_status_var.get(),
          f"aviso en la otra pestaña: {app.xl_status_var.get()}")
    check(app.tw_progress["maximum"] == 2, "la barra se ajusta al total")
    root.update()
    check("1 ronda de 2 trinos" in app.tw_status_var.get(),
          f"mensaje de arranque: {app.tw_status_var.get()}")
    check("Ollama" in app.tw_status_var.get(), "debe nombrar el motor")

    # Corpus arranca
    app.corpus_items = [{"id": "c1", "texto": "hola", "origen": "x"}]
    app.tweets = []
    registro.clear()
    app.start_corpus_coding()
    check(registro.get("arrancado"), "la codificación de corpus no arrancó")
    args = registro["args"]
    check(args[0] is app.corpus_items, "el worker debe recibir el corpus")
    check(args[5] is app.xl_progress, "barra de progreso del corpus")
    check(args[6] == app.export_corpus_excel, "exportador de corpus")
    check(args[8] == app.on_corpus_coding_cancelled, "callback de cancelación")
    check(estado(app.btn_xl_code) == "disabled", "btn_xl_code se apaga")
    check(estado(app.btn_xl_cancel) == "normal", "su cancelar se enciende")
    check(estado(app.btn_start) == "disabled", "la otra pestaña se bloquea")
    check("Hay un proceso activo en la pestaña Corpus" in app.tw_status_var.get(),
          f"aviso en la otra pestaña: {app.tw_status_var.get()}")
    root.update()
    check("1 ronda de 1 registros" in app.xl_status_var.get(),
          f"mensaje de arranque: {app.xl_status_var.get()}")

    # Con dos rondas el mensaje cambia de número
    registro.clear()
    RONDAS2 = [("ollama", "llama3.1:8b", None), ("claude", "sonnet", "low")]
    app._open_run_config = lambda modelo, n=0: RONDAS2
    app.start_corpus_coding()
    root.update()
    check("2 rondas" in app.xl_status_var.get(),
          f"con dos rondas: {app.xl_status_var.get()}")
    check("Ollama y Claude" in app.xl_status_var.get(),
          "debe nombrar los dos motores")
finally:
    habic.threading.Thread = hilo_real

app.tweets = []
app.corpus_items = []
app._process_running = False


# ══════════════════════════════════════════════════════════════════════════
seccion("Exportación a Excel (caracterización)")

app.codebook_var.set("bueno")
app.corpus_filepath = Path("mi_corpus.xlsx")
app.cuenta_var.set("@alguien")
app.desde_var.set("01/10/2023")
app.hasta_var.set("15/10/2023")

CATS2 = [{"codigo": "CRIT", "nombre": "Crítica"},
         {"codigo": "MOVI", "nombre": "Movilización"}]

RESULTS_TW = [
    {"item": {"id": "t1", "texto": "uno", "fecha": "01/10/2023",
              "cuenta": "alguien", "likes": "5", "reposts": "1",
              "respuestas": "0", "url": "http://x/1"},
     "coding": {"CRIT": {"valor": 1, "razon": "cita a"},
                "MOVI": {"valor": 0, "razon": "no b"}}},
    {"item": {"id": "t2", "texto": "dos", "fecha": "02/10/2023",
              "cuenta": "alguien", "likes": "0", "reposts": "0",
              "respuestas": "2", "url": "http://x/2"},
     "coding": {"CRIT": {"valor": "NC", "razon": "ambiguo"},
                "MOVI": {"valor": None, "razon": "sin respuesta"}}},
]
RESULTS_CO = [
    {"item": {"id": "c1", "texto": "uno", "origen": "diario"},
     "coding": {"CRIT": {"valor": 1, "razon": "cita a"},
                "MOVI": {"valor": 0, "razon": "no b"}}},
    {"item": {"id": "c2", "texto": "dos", "origen": "revista"},
     "coding": {"CRIT": {"valor": "NC", "razon": "ambiguo"},
                "MOVI": {"valor": None, "razon": "sin respuesta"}}},
]


def rejilla(path):
    """Todas las hojas como listas de listas, con la fecha normalizada."""
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        filas = []
        for r in range(1, ws.max_row + 1):
            fila = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and len(v) == 16 and v[4] == "-":
                    v = "<fecha>"          # la marca de tiempo cambia siempre
                fila.append(v)
            filas.append(fila)
        out[ws.title] = filas
    wb.close()
    return out


actual = {}

app._ronda_info = {"motor": "ollama", "modelo": "llama3.1:8b", "esfuerzo": None}
f_tw = app.export_twitter_excel(RESULTS_TW, CATS2, "llama3.1:8b", None)
f_co = app.export_corpus_excel(RESULTS_CO, CATS2, "llama3.1:8b", None)
actual["twitter_ollama"] = rejilla(TMP / f_tw)
actual["corpus_ollama"]  = rejilla(TMP / f_co)
check(f_tw.startswith("HABIC_TW_alguien_01-10-2023_15-10-2023_bueno"),
      f"nombre del archivo de Twitter: {f_tw}")
check(f_co.startswith("HABIC_CORPUS_mi_corpus_bueno"),
      f"nombre del archivo de corpus: {f_co}")

app._ronda_info = {"motor": "claude", "modelo": "sonnet", "esfuerzo": "low",
                   "coste": 1.2345}
f_co_cl = app.export_corpus_excel(RESULTS_CO, CATS2, "sonnet", 2)
actual["corpus_claude_ronda2"] = rejilla(TMP / f_co_cl)
check("_sonnet_c2" in f_co_cl, f"sufijo de ronda: {f_co_cl}")

app._ronda_info["parcial"] = True
f_parcial = app.export_corpus_excel(RESULTS_CO, CATS2, "sonnet", None)
actual["corpus_claude_parcial"] = rejilla(TMP / f_parcial)
check("_PARCIAL" in f_parcial, f"el parcial debe marcarse en el nombre: {f_parcial}")
app._ronda_info.pop("parcial")

# Comprobaciones explícitas sobre el contenido, además de la foto
co = actual["corpus_ollama"]["Corpus"]
check(co[0] == ["id", "texto", "origen", "CRIT", "MOVI", "tiene_codigo"],
      f"encabezados del corpus: {co[0]}")
check(co[1][3] == 1 and co[1][4] == 0, f"matriz mal escrita: {co[1]}")
check(co[2][3] == "N/C", f"el N/C debe verse en la matriz: {co[2]}")
check(co[2][4] is None, "el dato faltante debe quedar en celda vacía")
check(str(co[1][5]).startswith("=IF(SUM("), "falta la fórmula tiene_codigo")
tw = actual["twitter_ollama"]["Corpus"]
check(tw[0][:8] == ["id", "fecha", "cuenta", "texto", "likes", "reposts",
                    "respuestas", "url"], f"encabezados de Twitter: {tw[0]}")
raz = actual["corpus_ollama"]["Razones"]
check(raz[0] == ["id", "CRIT_razon", "MOVI_razon", "notas_codificador"],
      f"encabezados de Razones: {raz[0]}")
check(raz[1][1] == "cita a", "razón perdida")

met_ol = {f[0]: f[1] for f in actual["corpus_ollama"]["Método"] if f[0]}
check("seed 42" in met_ol["Parámetros de inferencia"], "Método de Ollama")
check(met_ol["Servicio externo"].startswith("No"), "Ollama es local")
met_cl = {f[0]: f[1] for f in actual["corpus_claude_ronda2"]["Método"] if f[0]}
check("NO es determinista" in met_cl["Parámetros de inferencia"],
      "el Método debe declarar que Claude no es determinista")
check(met_cl["Servicio externo"].startswith("Sí"), "Claude es externo")
check("1.2345" in met_cl["Coste de la ronda"], "el coste debe registrarse")
met_pa = {f[0]: f[1] for f in actual["corpus_claude_parcial"]["Método"] if f[0]}
check("PARCIAL" in met_pa["Estado"], "el parcial debe declararse en el Método")

# ── La foto: se crea la primera vez, se compara siempre después ──────────
if REFERENCIA.exists():
    esperado = json.loads(REFERENCIA.read_text(encoding="utf-8"))
    for nombre, hojas in esperado.items():
        check(nombre in actual, f"falta la exportación {nombre}")
        for hoja, filas in hojas.items():
            check(hoja in actual[nombre], f"falta la hoja {hoja} en {nombre}")
            check(actual[nombre][hoja] == filas,
                  f"CAMBIÓ el Excel en {nombre}/{hoja}:\n"
                  f"  antes: {filas}\n  ahora: {actual[nombre][hoja]}")
    print("Excel idéntico a la referencia guardada.")
else:
    REFERENCIA.write_text(json.dumps(actual, ensure_ascii=False, indent=1),
                          encoding="utf-8")
    print(f"Referencia creada: {REFERENCIA.name} "
          f"({len(actual)} exportaciones). A partir de ahora se compara.")


# ══════════════════════════════════════════════════════════════════════════
seccion("Muestra de validación y reporte de kappa")

app.val_data = ({"1": {"CRIT": 1}, "2": {"CRIT": 0}, "3": {"CRIT": 1}},
                ["CRIT"],
                {"1": {"texto": "uno"}, "2": {"texto": "dos"},
                 "3": {"texto": "tres"}},
                "origen.xlsx")
app.val_p_var.set("0.80")
app._recalc_validacion()
check("N = 3" in app.val_n_var.get(), f"recuento: {app.val_n_var.get()}")
n_val = habic.lacy_riffe_n(3, P=0.80)
app._generar_validacion()
generados = sorted(p.name for p in TMP.glob("HABIC_VALIDACION_*"))
check(len(generados) == 2, f"deben salir dos archivos: {generados}")
hum = [g for g in generados if "HUMANO" in g][0]
ia  = [g for g in generados if "_IA_" in g][0]
wb_h = openpyxl.load_workbook(TMP / hum)
wb_i = openpyxl.load_workbook(TMP / ia)
check(wb_h.active.cell(2, 3).value is None,
      "el archivo HUMANO debe ir vacío para codificar a ciegas")
check(wb_i.active.cell(2, 3).value in (0, 1),
      "el archivo IA debe traer los códigos del modelo")

# La muestra declara sus propios parámetros: sin esto, dentro de seis meses
# nadie sabe con qué P ni con qué nivel de confianza se sacó.
check("Método" in wb_h.sheetnames, f"falta la hoja Método: {wb_h.sheetnames}")
ws_m = wb_h["Método"]
met = {ws_m.cell(r, 1).value: ws_m.cell(r, 2).value
       for r in range(3, ws_m.max_row + 1) if ws_m.cell(r, 1).value}
check(met["Universo (N)"] == "3", f"N mal registrado: {met.get('Universo (N)')}")
check(met["Muestra (n)"] == str(n_val), f"n mal registrado: {met.get('Muestra (n)')}")
check(met["Acuerdo esperado (P)"] == "0.80", "P mal registrado")
check("z = 1,96" in met["Nivel de confianza"],
      f"debe declarar el nivel de confianza: {met.get('Nivel de confianza')}")
check("Lacy & Riffe" in met["Fórmula"], "debe declarar la fórmula")
check("Razones" not in met["Método"],
      "la muestra no tiene hoja de razones: no debe decir que la tiene")
wb_h.close(); wb_i.close()

# Con el estándar de dos colas la muestra es mayor que con una cola
check(habic.lacy_riffe_n(800) == 189,
      f"800 textos al 95% dos colas → 189: {habic.lacy_riffe_n(800)}")
check(habic.lacy_riffe_n(800, z=1.64) == 142, "una cola daría 142")

# Las tres cifras de κ: el agregado tapa a la categoría débil, el mínimo no
per_cat_demo = [
    {"cat": "FREC", "kappa": 0.80}, {"cat": "RARA", "kappa": 0.2331},
    {"cat": "SINDATOS", "kappa": None},
]
validos = [c for c in per_cat_demo if c["kappa"] is not None]
k_prom  = sum(c["kappa"] for c in validos) / len(validos)
peor    = min(validos, key=lambda c: c["kappa"])
check(abs(k_prom - 0.5166) < 0.001, f"promedio mal calculado: {k_prom}")
check(peor["cat"] == "RARA", "el mínimo debe señalar la categoría más débil")
check(habic.kappa_interpretation(peor["kappa"]) == "Leve",
      "la categoría débil debe salir etiquetada como tal")

k_res = {"k_global": 0.75, "po_global": 0.9, "pe_global": 0.6,
         "shared_ids": ["1", "2"], "shared_cats": ["CRIT"],
         "per_cat": [{"cat": "CRIT", "kappa": 0.75, "po": 0.9, "pe": 0.6,
                      "n_agree": 9, "n_total": 10}],
         "n_faltantes": 1}
desacuerdos = [{"id": "2", "categoria": "CRIT", "valor_a": 1, "valor_b": 0,
                "texto": "dos"}]
app.export_kappa_excel(k_res, desacuerdos)
kappas = sorted(TMP.glob("HABIC_KAPPA_*"))
check(kappas, "no se generó el reporte de kappa")
wb_k = openpyxl.load_workbook(kappas[-1])
check("Método" in wb_k.sheetnames, f"hojas del kappa: {wb_k.sheetnames}")
wb_k.close()


# ══════════════════════════════════════════════════════════════════════════
root.destroy()
print(f"\n{'=' * 62}")
print(f"OK — {_hechas[0]} comprobaciones pasadas.")
print(f"Archivos de prueba en: {TMP}")
print("No se tocó resultados/. Sin probar: scrape() y extract_tweet(), "
      "que necesitan Chrome con tu sesión de X.")
